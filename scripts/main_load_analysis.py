#!/usr/bin/env python3
"""Сводка загрузок GAS main: Supabase + WB API. Выход: xlsx."""

import csv
import json
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / f"main_load_analysis_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
LOG_EXPORTS = [
    ROOT / "docs" / "gas_main_log.tsv",
    ROOT / "docs" / "_Лог.tsv",
    ROOT / "docs" / "_log.tsv",
]
LOG_DONE = {"готово", "инкремент", "пропуск"}
LOG_PHASES = {"удаление", "сеть", "запись", "сортировка"}
SHEET_TO_DATASET = {}
for _ds, _sheets in {
    "products": ("Товары",),
    "product_prices": ("История цен",),
    "finance": ("Финансовые ежедневные",),
    "adv_stats": ("Рекламные расходы",),
    "sales_funnel": ("Воронка продаж",),
    "fbo_orders": ("Заказы FBW",),
    "fbo_sales": ("Продажи FBW",),
    "region_sales": ("Продажи по регионам",),
    "goods_returns": ("Возвраты товаров",),
    "order_feed": ("Лента заказов",),
    "feedbacks": ("Отзывы",),
    "questions": ("Вопросы",),
    "search_queries": ("Поисковые запросы",),
    "tariffs": ("Тарифы",),
    "wb_retentions": ("Удержания WB",),
    "warehouse_services": ("Складские услуги",),
    "claims": ("Заявки на возврат",),
    "promos": ("Акции",),
    "fbo_stocks": ("Остатки FBW",),
    "stocks_lk": ("Остатки ЛК",),
    "fbw_supplies": ("Поставки FBW",),
    "fbs_warehouses": ("Склады FBS",),
    "fbs_stocks": ("Остатки FBS",),
    "fbs_orders": ("Заказы FBS",),
    "fbs_supplies": ("Поставки FBS",),
}.items():
    for _sh in _sheets:
        SHEET_TO_DATASET[_sh] = _ds

# dataset -> (sheet, view, api)
META = {
    "products": ("Товары", "sheet_skus", "content-api POST /content/v2/get/cards/list + discounts-prices"),
    "product_prices": ("История цен", "sheet_product_prices_daily", "product_prices_daily ← products sync"),
    "finance": ("Финансовые ежедневные", "sheet_finance_rows", "finance-api POST /api/finance/v1/sales-reports/detailed"),
    "adv_stats": ("Рекламные расходы", "sheet_adv_stats", "advert-api GET /adv/v3/fullstats + /adv/v1/upd"),
    "sales_funnel": ("Воронка продаж", "sheet_sales_funnel_daily", "seller-analytics POST /api/analytics/v3/sales-funnel/products/history"),
    "fbo_orders": ("Заказы FBW", "sheet_fbo_orders", "statistics-api GET /api/v1/supplier/orders"),
    "fbo_sales": ("Продажи FBW", "sheet_fbo_sales", "statistics-api GET /api/v1/supplier/sales"),
    "region_sales": ("Продажи по регионам", "sheet_region_sales", "seller-analytics GET /api/v1/analytics/region-sale"),
    "goods_returns": ("Возвраты товаров", "sheet_goods_returns", "seller-analytics GET /api/v1/analytics/goods-return"),
    "order_feed": ("Лента заказов", "sheet_order_feed", "seller-analytics POST /api/analytics/v1/order-feed"),
    "feedbacks": ("Отзывы", "sheet_feedbacks", "feedbacks-api GET /api/v1/feedbacks"),
    "questions": ("Вопросы", "sheet_questions", "feedbacks-api GET /api/v1/questions"),
    "search_queries": ("Поисковые запросы", "sheet_search_queries", "seller-analytics POST /api/v2/search-report/product/search-texts (Jam)"),
    "tariffs": ("Тарифы", "sheet_tariffs", "common-api GET /api/v1/tariffs/* + supplies-api"),
    "wb_retentions": ("Удержания WB", "sheet_wb_retentions", "seller-analytics penalties/deductions/antifraud"),
    "warehouse_services": ("Складские услуги", "sheet_warehouse_services", "seller-analytics paid_storage + acceptance (async task)"),
    "claims": ("Заявки на возврат", "sheet_claims", "returns-api GET /api/v1/claims"),
    "promos": ("Акции", "sheet_promos", "dp-calendar-api GET /api/v1/calendar/promotions/*"),
    "fbo_stocks": ("Остатки FBW", "sheet_fbo_stocks", "seller-analytics GET /api/analytics/v1/stocks-report/wb-warehouses"),
    "stocks_lk": ("Остатки ЛК", "sheet_fbo_stocks_lk", "seller-analytics warehouse_remains task/download"),
    "fbw_supplies": ("Поставки FBW", "sheet_fbw", "supplies-api GET/POST /api/v1/supplies/*"),
    "fbs_warehouses": ("Склады FBS", "sheet_fbs_warehouses", "marketplace-api GET /api/v3/warehouses"),
    "fbs_offices": ("нет листа", "—", "marketplace-api GET /api/v3/offices"),
    "fbs_stocks": ("Остатки FBS", "sheet_fbs_stocks", "marketplace-api POST /api/v3/stocks/{warehouseId}"),
    "fbs_orders": ("Заказы FBS", "sheet_fbs_orders", "marketplace-api GET /api/v3/orders/* + status"),
    "fbs_supplies": ("Поставки FBS", "sheet_fbs", "marketplace-api GET /api/v3/supplies/*"),
}

ORDER = [
    "finance", "adv_stats", "sales_funnel", "products",
    "fbo_orders", "fbo_sales", "region_sales", "goods_returns", "order_feed",
    "feedbacks", "questions", "search_queries", "tariffs", "wb_retentions",
    "warehouse_services", "claims", "promos",
    "fbo_stocks", "stocks_lk", "fbw_supplies",
    "fbs_warehouses", "fbs_offices", "fbs_stocks", "fbs_orders", "fbs_supplies",
]

SQL = r"""
with seller as (
  select id from public.sellers where slug = 'main'
),
p as (
  select
    (timezone('Europe/Moscow', now()))::date as today,
    (timezone('Europe/Moscow', now()))::date - 1 as yesterday
),
lr as (
  select distinct on (lr.dataset)
    lr.dataset,
    lr.duration_ms,
    lr.rows_upserted,
    lr.rows_inserted,
    lr.bytes_inserted,
    lr.error,
    lr.started_at at time zone 'Europe/Moscow' as run_msk
  from public.load_runs lr
  join seller s on s.id = lr.seller_id
  where lr.mode = 'daily'
  order by lr.dataset, (lr.duration_ms is null), lr.started_at desc
),
daily_cnt as (
  select s.id as seller_id, p.today, p.yesterday from seller s, p p
),
dc as (
  select 'finance' as dataset, count(*)::bigint as n
  from public.finance_rows f, daily_cnt d
  where f.seller_id = d.seller_id and f.rr_dt = d.yesterday
  union all
  select 'adv_stats', count(*) from public.adv_stats a, daily_cnt d
  where a.seller_id = d.seller_id and a.stat_date = d.yesterday
  union all
  select 'fbo_orders', count(*) from public.fbo_orders o, daily_cnt d
  where o.seller_id = d.seller_id
    and (o.order_dt at time zone 'Europe/Moscow')::date = d.yesterday
  union all
  select 'fbo_sales', count(*) from public.fbo_sales s, daily_cnt d
  where s.seller_id = d.seller_id
    and (s.sale_dt at time zone 'Europe/Moscow')::date = d.yesterday
  union all
  select 'region_sales', count(*) from public.region_sales r, daily_cnt d
  where r.seller_id = d.seller_id and r.sale_dt = d.yesterday
  union all
  select 'goods_returns', count(*) from public.goods_returns g, daily_cnt d
  where g.seller_id = d.seller_id
    and (g.order_dt at time zone 'Europe/Moscow')::date = d.yesterday
  union all
  select 'order_feed', count(*) from public.order_feed f, daily_cnt d
  where f.seller_id = d.seller_id
    and (f.status_updated_at at time zone 'Europe/Moscow')::date = d.yesterday
  union all
  select 'feedbacks', count(*) from public.feedbacks f, daily_cnt d
  where f.seller_id = d.seller_id and f.feedback_dt = d.yesterday
  union all
  select 'questions', count(*) from public.questions q, daily_cnt d
  where q.seller_id = d.seller_id and q.question_dt = d.yesterday
  union all
  select 'sales_funnel', count(*) from public.sales_funnel_daily s, daily_cnt d
  where s.seller_id = d.seller_id and s.stat_date = d.yesterday
  union all
  select 'fbo_stocks', count(*) from public.fbo_stocks s, daily_cnt d
  where s.seller_id = d.seller_id and s.snapshot_date = d.today
  union all
  select 'stocks_lk', count(*) from public.fbo_stocks_lk s, daily_cnt d
  where s.seller_id = d.seller_id and s.snapshot_date = d.today
  union all
  select 'fbs_stocks', count(*) from public.fbs_stocks s, daily_cnt d
  where s.seller_id = d.seller_id and s.snapshot_date = d.today
  union all
  select 'claims', count(*) from public.claims c, daily_cnt d
  where c.seller_id = d.seller_id
    and (c.claim_dt at time zone 'Europe/Moscow')::date = d.today
  union all
  select 'promos', count(*) from public.promos p, daily_cnt d
  where p.seller_id = d.seller_id
    and p.start_dt <= d.today and p.end_dt >= d.today
  union all
  select 'tariffs', count(*) from public.sheet_tariffs t, daily_cnt d
  where t.seller_id = d.seller_id
  union all
  select 'warehouse_services', count(*) from public.warehouse_services w, daily_cnt d
  where w.seller_id = d.seller_id and w.service_dt = d.yesterday
  union all
  select 'wb_retentions', count(*) from public.wb_retentions w, daily_cnt d
  where w.seller_id = d.seller_id and w.retention_dt = d.yesterday
  union all
  select 'products', count(*) from public.products pr, daily_cnt d
  where pr.seller_id = d.seller_id
    and (pr.loaded_at at time zone 'Europe/Moscow')::date = d.today
  union all
  select 'fbw_supplies', count(*) from public.fbw_supplies f, daily_cnt d
  where f.seller_id = d.seller_id
    and (f.loaded_at at time zone 'Europe/Moscow')::date = d.today
  union all
  select 'fbs_orders', count(*) from public.fbs_orders o, daily_cnt d
  where o.seller_id = d.seller_id
    and (o.created_at at time zone 'Europe/Moscow')::date = d.yesterday
  union all
  select 'fbs_supplies', count(*) from public.fbs_supplies s, daily_cnt d
  where s.seller_id = d.seller_id
    and (s.loaded_at at time zone 'Europe/Moscow')::date = d.today
  union all
  select 'fbs_warehouses', count(*) from public.fbs_warehouses w, daily_cnt d
  where w.seller_id = d.seller_id
    and (w.loaded_at at time zone 'Europe/Moscow')::date = d.today
  union all
  select 'fbs_offices', count(*) from public.fbs_offices o, daily_cnt d
  where o.seller_id = d.seller_id
    and (o.loaded_at at time zone 'Europe/Moscow')::date = d.today
),
disk as (
  select
    d.dataset,
    tbl as table_name,
    case when tbl is null then null::bigint else pg_relation_size(rel) end as heap_bytes,
    case when tbl is null then null::bigint else pg_indexes_size(rel) end as index_bytes,
    case
      when tbl is null then null::bigint
      else greatest(pg_total_relation_size(rel) - pg_relation_size(rel) - pg_indexes_size(rel), 0)
    end as toast_bytes,
    case when tbl is null then null::bigint else pg_total_relation_size(rel) end as disk_total_bytes
  from unnest(public.daily_datasets()) as d(dataset)
  cross join lateral (
    select public.dataset_primary_table(d.dataset) as tbl
  ) x
  cross join lateral (
    select case when tbl is null then null::regclass else ('public.' || tbl)::regclass end as rel
  ) r
)
select
  st.dataset,
  public.dataset_primary_table(st.dataset) as table_name,
  public.dataset_access_enabled(st.seller_id, st.dataset) as access_enabled,
  st.row_count,
  st.row_bytes,
  dk.heap_bytes,
  dk.index_bytes,
  dk.toast_bytes,
  dk.disk_total_bytes,
  lr.duration_ms,
  lr.rows_upserted,
  lr.rows_inserted,
  lr.bytes_inserted,
  lr.error as last_error,
  lr.run_msk,
  coalesce(dc.n, 0) as daily_rows_data,
  g.rows_written as gas_rows,
  g.duration_ms as gas_duration_ms,
  g.updated_at at time zone 'Europe/Moscow' as gas_updated_msk
from public.sync_state st
join seller s on s.id = st.seller_id
left join lr on lr.dataset = st.dataset
left join dc on dc.dataset = st.dataset
left join disk dk on dk.dataset = st.dataset
left join public.gas_sheet_updates g on g.dataset = st.dataset
where st.dataset = any(public.daily_datasets())
order by st.dataset;
"""

SQL_EXTRA_TABLES = r"""
with daily_tbl as (
  select distinct public.dataset_primary_table(d.dataset) as table_name
  from unnest(public.daily_datasets()) as d(dataset)
  where public.dataset_primary_table(d.dataset) is not null
)
select
  c.relname as table_name,
  pg_relation_size(c.oid) as heap_bytes,
  pg_indexes_size(c.oid) as index_bytes,
  greatest(
    pg_total_relation_size(c.oid) - pg_relation_size(c.oid) - pg_indexes_size(c.oid),
    0
  ) as toast_bytes,
  pg_total_relation_size(c.oid) as disk_total_bytes
from pg_class c
join pg_namespace n on n.oid = c.relnamespace
where n.nspname = 'public'
  and c.relkind = 'r'
  and c.relname not in (select table_name from daily_tbl)
order by pg_total_relation_size(c.oid) desc;
"""


def db_rpc(name: str, body: dict) -> dict | None:
    payload = json.dumps(body, ensure_ascii=False)
    sql = f"select public.{name}($${payload}$$::jsonb) as result"
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".sql", delete=False, encoding="utf-8"
    ) as f:
        f.write(sql)
        sql_path = f.name
    proc = subprocess.run(
        ["npx", "supabase", "db", "query", "--linked", "-f", sql_path],
        cwd=ROOT,
        capture_output=True,
        text=True,
        shell=True,
    )
    Path(sql_path).unlink(missing_ok=True)
    if proc.returncode != 0:
        return None
    raw = proc.stdout.strip()
    start = raw.find("{")
    if start < 0:
        return None
    data = json.loads(raw[start:])
    rows = data.get("rows") or []
    if not rows:
        return None
    return rows[0].get("result")


def parse_log_day(ts: str) -> str:
    ts = (ts or "").strip()
    if len(ts) >= 10 and ts[2] == "." and ts[5] == ".":
        return f"{ts[6:10]}-{ts[3:5]}-{ts[0:2]}"
    if len(ts) >= 10 and ts[4] == "-":
        return ts[:10]
    return ""


def parse_log_rows(path: Path) -> dict[str, dict]:
    """dataset -> {duration_ms, rows, sheet} из экспорта _Лог (сумма фаз за день)."""
    today = datetime.now().strftime("%Y-%m-%d")
    by_sheet: dict[str, dict] = {}
    with path.open(encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        dialect = csv.excel_tab if "\t" in sample else csv.excel
        reader = csv.reader(f, dialect)
        header = next(reader, None)
        if not header:
            return {}
        h = {str(x).strip().lower(): i for i, x in enumerate(header)}
        i_time = h.get("время", 0)
        i_sheet = h.get("лист", 1)
        i_status = h.get("статус", 2)
        i_rows = h.get("строк", 3)
        i_ms = h.get("длительность_мс", 7)
        for row in reader:
            if len(row) <= max(i_time, i_sheet, i_status, i_ms):
                continue
            sheet = str(row[i_sheet] or "").strip()
            status = str(row[i_status] or "").strip()
            if not sheet:
                continue
            if parse_log_day(str(row[i_time])) != today:
                continue
            ms = int(float(row[i_ms] or 0))
            if ms <= 0:
                continue
            agg = by_sheet.setdefault(
                sheet, {"phase_sum": 0, "done_ms": 0, "done_rows": 0, "done_ts": ""}
            )
            if status in LOG_PHASES:
                agg["phase_sum"] += ms
            elif status in LOG_DONE:
                ts = str(row[i_time] or "")
                if ts >= agg["done_ts"]:
                    agg["done_ts"] = ts
                    agg["done_ms"] = ms
                    agg["done_rows"] = int(float(row[i_rows] or 0))
    out: dict[str, dict] = {}
    for sheet, agg in by_sheet.items():
        ds = SHEET_TO_DATASET.get(sheet)
        if not ds:
            continue
        total = agg["phase_sum"] if agg["phase_sum"] > 0 else agg["done_ms"]
        if total <= 0:
            continue
        out[ds] = {
            "dataset": ds,
            "sheet_name": sheet,
            "rows": agg["done_rows"],
            "duration_ms": total,
        }
    return out


def sync_gas_durations_from_log_export() -> int:
    for path in LOG_EXPORTS:
        if not path.is_file():
            continue
        items = list(parse_log_rows(path).values())
        if not items:
            continue
        res = db_rpc("gas_sheet_sync_durations", {"p_items": items})
        if res and res.get("updated"):
            print(f"GAS время из {path.name}: {res['updated']} датасетов")
            return int(res["updated"])
        n = 0
        for item in items:
            one = db_rpc(
                "gas_sheet_touch",
                {
                    "p_dataset": item["dataset"],
                    "p_sheet_name": item["sheet_name"],
                    "p_rows": item["rows"],
                    "p_duration_ms": item["duration_ms"],
                },
            )
            if one is not None:
                n += 1
        if n:
            print(f"GAS время из {path.name}: {n} датасетов")
            return n
    return 0


def format_db_error(text: str) -> str:
    t = text or ""
    if "password authentication failed" in t or "SUPABASE_DB_PASSWORD" in t:
        return (
            "Не удалось подключиться к Supabase DB.\n"
            "  1. В PowerShell: npx supabase login\n"
            "  2. В папке wb-supabase: npx supabase link\n"
            "  3. Или задайте пароль: $env:SUPABASE_DB_PASSWORD = \"...\""
        )
    if "not found" in t.lower() and "npx" in t.lower():
        return "Не найден npx. Установите Node.js: https://nodejs.org"
    if "Initialising login role" in t and not t.strip().endswith("}"):
        return "Supabase CLI завис или оборвал соединение. Повторите через минуту."
    return t.strip()[-2000:] or "неизвестная ошибка Supabase CLI"


def db_query(sql: str) -> list[dict]:
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".sql", delete=False, encoding="utf-8"
    ) as f:
        f.write(sql)
        sql_path = f.name
    proc = subprocess.run(
        ["npx", "supabase", "db", "query", "--linked", "-f", sql_path],
        cwd=ROOT,
        capture_output=True,
        text=True,
        shell=True,
    )
    Path(sql_path).unlink(missing_ok=True)
    merged = (proc.stdout or "") + "\n" + (proc.stderr or "")
    if proc.returncode != 0:
        raise RuntimeError(format_db_error(merged))
    raw = proc.stdout.strip()
    start = raw.find("{")
    if start < 0:
        raise RuntimeError(format_db_error(merged))
    payload = json.loads(raw[start:])
    if payload.get("_tag") == "Error":
        raise RuntimeError(format_db_error(json.dumps(payload, ensure_ascii=False)))
    return payload.get("rows") or []


def yes_no(flag) -> str:
    if flag is True:
        return "да"
    if flag is False:
        return "нет"
    return ""


def fmt_mb(b: int | None) -> float | None:
    if b is None:
        return None
    return round(b / 1024.0 / 1024.0, 2)


def fmt_min(ms: int | None) -> float | None:
    if ms is None:
        return None
    return round(ms / 60000.0, 2)


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="434343")
    font = Font(bold=True, color="FFFFFF", size=10)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, max_w=48):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(10, min(max(len(str(c.value or "")) for c in col) + 2, max_w))
        ws.column_dimensions[letter].width = width


def sum_mb(rows: list[dict], key: str) -> float:
    return round(sum(r.get(key) or 0 for r in rows) / 1024.0 / 1024.0, 2)


def load_log_durations() -> dict[str, dict]:
    merged: dict[str, dict] = {}
    for path in LOG_EXPORTS:
        if path.is_file():
            merged.update(parse_log_rows(path))
    return merged


def main():
    print(f"Папка проекта: {ROOT}", flush=True)
    print("Запрос к Supabase (обычно 1–2 мин, не прерывайте)...", flush=True)
    sync_gas_durations_from_log_export()
    log_durations = load_log_durations()
    rows = db_query(SQL)
    print("Основной запрос готов, прочие таблицы...", flush=True)
    extra = db_query(SQL_EXTRA_TABLES)
    by_ds = {r["dataset"]: r for r in rows}

    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    headers = [
        "Лист GAS main",
        "Код датасета",
        "Таблица Supabase",
        "View GAS",
        "API WB",
        "Доступ WB",
        "БД: мин (WB→Supabase)",
        "GAS: мин (из _Лог→БД)",
        "Всего строк в БД",
        "Строк за день (данные)",
        "rows_inserted daily",
        "rows_upserted daily",
        "Данные МБ (pg_column_size)",
        "Heap МБ",
        "Индексы МБ",
        "TOAST МБ",
        "Диск всего МБ",
        "Прирост МБ за daily",
        "Строк в листе GAS",
        "GAS обновлён МСК",
        "Daily run МСК",
        "Последняя ошибка",
    ]
    ws.append(headers)
    style_header(ws)

    totals = {
        "rows": 0,
        "daily": 0,
        "data_mb": 0.0,
        "heap_mb": 0.0,
        "idx_mb": 0.0,
        "toast_mb": 0.0,
        "disk_mb": 0.0,
        "mb_d": 0.0,
        "db_min": 0.0,
        "gas_min": 0.0,
    }

    for ds in ORDER:
        m = META.get(ds, ("?", "?", "?"))
        r = by_ds.get(ds, {})
        sheet, view, api = m
        db_min = fmt_min(r.get("duration_ms"))
        gas_ms = r.get("gas_duration_ms") or (log_durations.get(ds) or {}).get("duration_ms")
        gas_min = fmt_min(gas_ms)
        data_mb = fmt_mb(r.get("row_bytes"))
        heap_mb = fmt_mb(r.get("heap_bytes"))
        idx_mb = fmt_mb(r.get("index_bytes"))
        toast_mb = fmt_mb(r.get("toast_bytes"))
        disk_mb = fmt_mb(r.get("disk_total_bytes"))
        mb_d = fmt_mb(r.get("bytes_inserted"))
        daily_data = r.get("daily_rows_data") or 0
        row_cnt = r.get("row_count") or 0

        ws.append([
            sheet,
            ds,
            r.get("table_name") or "",
            view,
            api,
            yes_no(r.get("access_enabled")),
            db_min,
            gas_min,
            row_cnt,
            daily_data,
            r.get("rows_inserted"),
            r.get("rows_upserted"),
            data_mb,
            heap_mb,
            idx_mb,
            toast_mb,
            disk_mb,
            mb_d,
            r.get("gas_rows"),
            r.get("gas_updated_msk"),
            r.get("run_msk"),
            (r.get("last_error") or "")[:120],
        ])
        totals["rows"] += row_cnt or 0
        totals["daily"] += daily_data or 0
        totals["data_mb"] += data_mb or 0
        totals["heap_mb"] += heap_mb or 0
        totals["idx_mb"] += idx_mb or 0
        totals["toast_mb"] += toast_mb or 0
        totals["disk_mb"] += disk_mb or 0
        totals["mb_d"] += mb_d or 0
        totals["db_min"] += db_min or 0
        totals["gas_min"] += gas_min or 0

    ws.append([])
    ws.append([
        "ИТОГО daily", "", "", "", "", "",
        round(totals["db_min"], 2),
        round(totals["gas_min"], 2),
        totals["rows"],
        totals["daily"],
        "", "",
        round(totals["data_mb"], 1),
        round(totals["heap_mb"], 1),
        round(totals["idx_mb"], 1),
        round(totals["toast_mb"], 1),
        round(totals["disk_mb"], 1),
        round(totals["mb_d"], 2),
        "", "", "", "",
    ])

    # Прочие таблицы (не daily-датасеты)
    ws5 = wb.create_sheet("Прочие таблицы")
    ws5.append(["Таблица", "Heap МБ", "Индексы МБ", "TOAST МБ", "Диск всего МБ"])
    style_header(ws5)
    extra_totals = {"heap": 0, "idx": 0, "toast": 0, "disk": 0}
    for r in extra:
        h = r.get("heap_bytes") or 0
        i = r.get("index_bytes") or 0
        t = r.get("toast_bytes") or 0
        d = r.get("disk_total_bytes") or 0
        extra_totals["heap"] += h
        extra_totals["idx"] += i
        extra_totals["toast"] += t
        extra_totals["disk"] += d
        ws5.append([
            r.get("table_name"),
            fmt_mb(h),
            fmt_mb(i),
            fmt_mb(t),
            fmt_mb(d),
        ])
    ws5.append([])
    ws5.append([
        "ИТОГО прочие",
        fmt_mb(extra_totals["heap"]),
        fmt_mb(extra_totals["idx"]),
        fmt_mb(extra_totals["toast"]),
        fmt_mb(extra_totals["disk"]),
    ])
    ws5.append([
        "ИТОГО проект (daily + прочие)",
        round(totals["heap_mb"] + (extra_totals["heap"] / 1024 / 1024), 1),
        round(totals["idx_mb"] + (extra_totals["idx"] / 1024 / 1024), 1),
        round(totals["toast_mb"] + (extra_totals["toast"] / 1024 / 1024), 1),
        round(totals["disk_mb"] + (extra_totals["disk"] / 1024 / 1024), 1),
    ])
    autosize(ws5)

    ws2 = wb.create_sheet("Топ БД мин")
    ws2.append(["Датасет", "Лист", "БД мин", "Строк в БД", "Диск МБ"])
    style_header(ws2)
    ranked = sorted(
        [(ds, by_ds.get(ds, {})) for ds in ORDER],
        key=lambda x: x[1].get("duration_ms") or 0,
        reverse=True,
    )
    for ds, r in ranked:
        if not r.get("duration_ms"):
            continue
        m = META[ds]
        ws2.append([
            ds,
            m[0],
            fmt_min(r.get("duration_ms")),
            r.get("row_count"),
            fmt_mb(r.get("disk_total_bytes")),
        ])
    autosize(ws2)

    ws3 = wb.create_sheet("Топ GAS мин")
    ws3.append(["Датасет", "Лист", "GAS мин", "Строк в листе"])
    style_header(ws3)
    ranked = sorted(
        [(ds, by_ds.get(ds, {})) for ds in ORDER],
        key=lambda x: (
            x[1].get("gas_duration_ms")
            or (log_durations.get(x[0]) or {}).get("duration_ms")
            or 0
        ),
        reverse=True,
    )
    for ds, r in ranked:
        gas_ms = r.get("gas_duration_ms") or (log_durations.get(ds) or {}).get("duration_ms")
        if not gas_ms:
            continue
        m = META[ds]
        ws3.append([ds, m[0], fmt_min(gas_ms), r.get("gas_rows")])
    autosize(ws3)

    ws4 = wb.create_sheet("Методология")
    lines = [
        "Отчёт: GAS main, кабинет slug=main",
        f"Сформирован: {datetime.now().strftime('%Y-%m-%d %H:%M')} (локальное время скрипта)",
        "",
        "Доступ WB — public.dataset_access_enabled (= list_dataset_access.enabled).",
        "  Для main по умолчанию «да», пока нет явного enabled=false в sync_dataset_access.",
        "  Это доступ к WB API / sync-tick, не колонка «Включено» в _Настройки книги.",
        "",
        "БД: мин — load_runs.duration_ms, последний daily run с duration (не только сегодня).",
        "GAS: мин — сумма фаз (удаление+сеть+запись+сортировка) за день из _Лог; иначе «готово».",
        "  Источники (по приоритету): gas_sheet_updates.duration_ms; экспорт _Лог → docs/gas_main_log.tsv.",
        "  Заполнить БД: GAS main → Сервис → «Время GAS → база (из _Лог)» (один клик).",
        "",
        "Данные МБ — sync_state.row_bytes: SUM(pg_column_size) строк seller=main (логический вес).",
        "Heap / Индексы / TOAST / Диск всего — физический размер таблицы целиком (все seller_id).",
        "  Индексы — ускоряют запросы (seller_id, даты); убирать нельзя без потери скорости.",
        "  TOAST — автоматическое хранение длинных полей PostgreSQL, не настраивается отдельно.",
        "  Диск всего ≈ то, что видно в Supabase Table Editor / Database size.",
        "Лист «Прочие таблицы» — load_runs, finance_kvv, info_facts_daily и т.д.",
    ]
    for i, line in enumerate(lines, 1):
        ws4.cell(i, 1, line)
    ws4.column_dimensions["A"].width = 110

    autosize(ws)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nПрервано.", file=sys.stderr)
        sys.exit(130)
    except Exception as e:
        print(f"\nОшибка: {e}", file=sys.stderr)
        print(
            "\nКак запустить в Windows (из папки wb-supabase):\n"
            "  py -3 scripts\\main_load_analysis.py\n"
            "  или двойной клик: scripts\\run_main_load_analysis.cmd\n"
            "\nНе используйте python3 — на Windows часто не работает.",
            file=sys.stderr,
        )
        sys.exit(1)
