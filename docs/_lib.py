# -*- coding: utf-8 -*-
"""Общие стили для Word-инструкций."""

import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu, Twips

NAVY = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x2D, 0x2D, 0x2D)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1B, 0x7A, 0x3D)
RED = RGBColor(0xB3, 0x26, 0x1E)
TEAL = RGBColor(0x0F, 0x6B, 0x6B)

HEADER_BG = "1F4E79"
ALT_BG = "F2F2F2"
NOTE_BG = "FFF2CC"
OK_BG = "E7F3EA"
EXTRA_BG = "E3F2FD"
BAD_BG = "F8E6E4"
SHEET_HDR = "F1F1F1"
SHEET_TOT = "E8E8E8"
SHEET_TAB = "3D3D3D"
CARD_BG = "F7F7F7"


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            el = tcBorders.find(qn("w:" + edge))
            if el is None:
                el = OxmlElement("w:" + edge)
                tcBorders.append(el)
            for k, v in kwargs[edge].items():
                el.set(qn("w:" + k), str(v))


def font(run, size=11, bold=False, color=DARK, name="Calibri"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)
    rFonts.set(qn("w:eastAsia"), name)


def pset(p, after=6, before=0, align=None):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align:
        p.alignment = align


def new_doc(landscape=False):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.6)
        if landscape:
            section.page_width, section.page_height = section.page_height, section.page_width
    return doc


def add_cover(doc, kicker, title, subtitle, date="01.09.2026"):
    bar = doc.add_table(rows=1, cols=1)
    bar.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = bar.cell(0, 0)
    shade(cell, HEADER_BG)
    p = cell.paragraphs[0]
    pset(p, after=0, before=0)
    r = p.add_run(kicker)
    font(r, 11, True, WHITE)
    p2 = cell.add_paragraph()
    pset(p2, after=2, before=8)
    r2 = p2.add_run(title)
    font(r2, 22, True, WHITE)
    p3 = cell.add_paragraph()
    pset(p3, after=4, before=4)
    r3 = p3.add_run(subtitle)
    font(r3, 12, False, WHITE)
    p4 = cell.add_paragraph()
    pset(p4, after=0, before=2)
    r4 = p4.add_run(date)
    font(r4, 10, False, WHITE)
    doc.add_paragraph()


def add_title(doc, text):
    p = doc.add_paragraph()
    pset(p, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    font(r, 20, True, NAVY)


def add_h(doc, text):
    p = doc.add_paragraph()
    pset(p, after=8, before=16)
    r = p.add_run(text)
    font(r, 15, True, NAVY)


def add_h2(doc, text):
    p = doc.add_paragraph()
    pset(p, after=6, before=10)
    r = p.add_run(text)
    font(r, 12, True, NAVY)


def add_body(doc, text, size=11):
    p = doc.add_paragraph()
    pset(p, after=8)
    r = p.add_run(text)
    font(r, size, False, DARK)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Calibri")
    rFonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rFonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rPr.append(sz)
    new_run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.set(qn("xml:space"), "preserve")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_book_link(doc, url):
    p = doc.add_paragraph()
    pset(p, after=8)
    r = p.add_run("Книга: ")
    font(r, 11, False, DARK)
    add_hyperlink(p, "открыть в Google", url)


def add_note(doc, text, bg=NOTE_BG):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, bg)
    p = cell.paragraphs[0]
    pset(p, after=0)
    r = p.add_run(text)
    font(r, 11, False, DARK)
    doc.add_paragraph()


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.clear()
    r = p.add_run(text)
    font(r, 11, False, DARK)


def add_step(doc, n, title, text):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    c0, c1 = t.rows[0].cells
    shade(c0, HEADER_BG)
    p0 = c0.paragraphs[0]
    pset(p0, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r0 = p0.add_run(str(n))
    font(r0, 16, True, WHITE)
    shade(c1, CARD_BG)
    p1 = c1.paragraphs[0]
    pset(p1, after=2)
    r1 = p1.add_run(title)
    font(r1, 12, True, NAVY)
    p2 = c1.add_paragraph()
    pset(p2, after=0)
    r2 = p2.add_run(text)
    font(r2, 11, False, DARK)
    c0.width = Cm(1.2)
    c1.width = Cm(16.5)
    doc.add_paragraph()


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        shade(cell, HEADER_BG)
        p = cell.paragraphs[0]
        pset(p, after=0)
        r = p.add_run(h)
        font(r, 10, True, WHITE)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            if ri % 2 == 1:
                shade(cell, ALT_BG)
            p = cell.paragraphs[0]
            pset(p, after=0)
            r = p.add_run(str(val))
            font(r, 10, False, DARK)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def flow(doc, boxes):
    """Горизонтальная схема: [A] → [B] → [C]."""
    cols = len(boxes) * 2 - 1
    t = doc.add_table(rows=1, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    ci = 0
    for i, box in enumerate(boxes):
        cell = t.rows[0].cells[ci]
        shade(cell, box.get("bg", HEADER_BG))
        p = cell.paragraphs[0]
        pset(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(box["title"])
        font(r, 10, True, WHITE if box.get("bg", HEADER_BG) in (HEADER_BG, SHEET_TAB, "0F6B6B") else DARK)
        if box.get("sub"):
            p2 = cell.add_paragraph()
            pset(p2, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            r2 = p2.add_run(box["sub"])
            font(r2, 8, False, WHITE if box.get("bg", HEADER_BG) in (HEADER_BG, SHEET_TAB, "0F6B6B") else GRAY)
        ci += 1
        if i < len(boxes) - 1:
            arrow = t.rows[0].cells[ci]
            p = arrow.paragraphs[0]
            pset(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run("→")
            font(r, 16, True, NAVY)
            ci += 1
    doc.add_paragraph()


def sheet_mock(doc, tab, caption, headers, rows, total=None, a1a2=None):
    """Макет листа Google: вкладка + шапка + даты A1/A2 + таблица как в книге."""
    wrap = doc.add_table(rows=1, cols=1)
    wrap.alignment = WD_TABLE_ALIGNMENT.CENTER
    outer = wrap.cell(0, 0)
    shade(outer, "FAFAFA")
    outer.paragraphs[0].clear()
    tab_t = outer.add_table(rows=1, cols=3)
    tab_t.rows[0].cells[0].text = ""
    shade(tab_t.rows[0].cells[0], SHEET_TAB)
    ptab = tab_t.rows[0].cells[0].paragraphs[0]
    ptab.clear()
    pset(ptab, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    rt = ptab.add_run(tab)
    font(rt, 9, True, WHITE)
    shade(tab_t.rows[0].cells[1], "E0E0E0")
    shade(tab_t.rows[0].cells[2], "E0E0E0")
    pcap = outer.add_paragraph()
    pset(pcap, after=4, before=6)
    rc = pcap.add_run(caption)
    font(rc, 10, False, GRAY)
    if a1a2:
        meta = outer.add_paragraph()
        pset(meta, after=4)
        rm = meta.add_run("A1  %s     A2  %s" % a1a2)
        font(rm, 9, True, NAVY)
    inner = outer.add_table(rows=1 + len(rows) + (1 if total else 0), cols=len(headers))
    inner.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = inner.rows[0].cells[i]
        shade(cell, SHEET_HDR)
        p = cell.paragraphs[0]
        pset(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(h)
        font(r, 8, True, DARK)
    start = 1
    if total:
        for ci, val in enumerate(total):
            cell = inner.rows[1].cells[ci]
            shade(cell, SHEET_TOT)
            p = cell.paragraphs[0]
            pset(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER if ci >= 2 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(str(val))
            font(r, 8, True, DARK)
        start = 2
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = inner.rows[start + ri].cells[ci]
            p = cell.paragraphs[0]
            pset(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER if ci >= 2 else WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(str(val))
            font(r, 8, False, DARK)
    doc.add_paragraph()


def kpi_row(doc, items):
    t = doc.add_table(rows=2, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, it in enumerate(items):
        c0 = t.rows[0].cells[i]
        c1 = t.rows[1].cells[i]
        shade(c0, it.get("bg", HEADER_BG))
        p = c0.paragraphs[0]
        pset(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(it["label"])
        font(r, 8, False, WHITE)
        shade(c1, CARD_BG)
        p1 = c1.paragraphs[0]
        pset(p1, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
        r1 = p1.add_run(it["value"])
        font(r1, 14, True, NAVY)
    doc.add_paragraph()


TG_BG = "E8F3D6"
TG_WHO = RGBColor(0x4E, 0x7A, 0x32)
TG_LINE = RGBColor(0x2C, 0x2C, 0x2C)
TG_MUTED = RGBColor(0x6B, 0x6B, 0x6B)
TG_UP = RGBColor(0x1B, 0x7A, 0x3D)
TG_DOWN = RGBColor(0xC0, 0x39, 0x2B)
TG_OK = RGBColor(0x1B, 0x7A, 0x3D)
TG_HEADER_MARKS = ("💰", "📢", "🚚", "🏬", "📦", "📥")
TG_PCT_RE = re.compile(r"^[+\-−]?\s*\d+[.,]?\d*\s*%$")

TG_DIGEST = [
    "💰 Финансы",
    "Продажи  78 шт  +3%",
    "Выручка  20 258 ₽  −7%",
    "К перечисл.  13 390 ₽  −5%",
    "",
    "📢 Реклама",
    "Расход  3 844 ₽  +19%",
    "ДРР  19,0%",
    "После рекламы  9 545 ₽  −12%",
    "",
    "🚚 Заказы FBW",
    "Заказы  85  +2%",
    "Отмены  2",
    "Сумма  22 884 ₽",
    "",
    "🏬 Заказы FBS",
    "Заказы  14  −18%",
    "Отмены  1",
    "Сумма  3 093 ₽",
    "",
    "📦 Остатки",
    "FBW  6 192  0%",
    "FBS  3 503",
    "В пути  606",
    "",
    "📥 Поставки",
    "Приёмка  1 · не запланировано  2",
    "FBS открытые  1",
    "",
    "✅ База 11/11",
]


def _tbl_width(table, cm):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(cm * 567)))
    tblW.set(qn("w:type"), "dxa")


def _tbl_no_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            borders.append(el)
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")


def _cell_margins(cell, top, left, bottom, right):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = tcMar.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            tcMar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


def _pct_color(pct):
    s = (pct or "").strip().replace("−", "-")
    if not s:
        return TG_MUTED
    if s.startswith("+"):
        return TG_UP
    if s.startswith("-"):
        return TG_DOWN
    return TG_MUTED


def _split_emoji(text):
    for mark in TG_HEADER_MARKS + ("✅", "❌"):
        if text.startswith(mark):
            rest = text[len(mark):]
            if rest.startswith(" "):
                rest = rest[1:]
            return mark, rest
    return "", text


def _is_header(line):
    s = line.strip()
    return any(s.startswith(m) for m in TG_HEADER_MARKS)


def _is_status(line):
    s = line.strip()
    return s.startswith("✅") or s.startswith("❌")


def _split_metric(line):
    parts = [p for p in re.split(r"\s{2,}", line.strip()) if p]
    if len(parts) < 2:
        return None
    if len(parts) >= 3 and TG_PCT_RE.match(parts[-1]):
        return parts[0], "  ".join(parts[1:-1]), parts[-1]
    return parts[0], "  ".join(parts[1:]), ""


def _tg_run(p, text, size=10.5, bold=False, color=TG_LINE, name="Segoe UI", emoji=False):
    r = p.add_run(text)
    font(r, size, bold, color, "Segoe UI Emoji" if emoji else name)
    return r


def _add_metric_table(cell, rows):
    inner = cell.add_table(rows=len(rows), cols=3)
    _tbl_no_borders(inner)
    _tbl_width(inner, 8.4)
    for i, (label, value, pct) in enumerate(rows):
        c0, c1, c2 = inner.rows[i].cells
        c0.width = Cm(3.5)
        c1.width = Cm(3.3)
        c2.width = Cm(1.6)
        _cell_margins(c0, 20, 0, 20, 40)
        _cell_margins(c1, 20, 40, 20, 40)
        _cell_margins(c2, 20, 0, 20, 0)
        p0 = c0.paragraphs[0]
        pset(p0, after=0)
        _tg_run(p0, label, 10.5, False, TG_LINE)
        p1 = c1.paragraphs[0]
        pset(p1, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _tg_run(p1, value, 10.5, True, TG_LINE)
        p2 = c2.paragraphs[0]
        pset(p2, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if pct:
            _tg_run(p2, pct, 10.5, True, _pct_color(pct))


def add_telegram(doc, title, lines):
    wrap = doc.add_table(rows=1, cols=1)
    wrap.alignment = WD_TABLE_ALIGNMENT.LEFT
    _tbl_width(wrap, 9.4)
    _tbl_no_borders(wrap)
    cell = wrap.cell(0, 0)
    shade(cell, TG_BG)
    _cell_margins(cell, 80, 140, 100, 140)
    set_cell_border(
        cell,
        top={"sz": "4", "val": "single", "color": "C5DCB0"},
        left={"sz": "24", "val": "single", "color": "67A83A"},
        bottom={"sz": "8", "val": "single", "color": "C5DCB0"},
        right={"sz": "4", "val": "single", "color": "C5DCB0"},
    )
    cell.paragraphs[0].clear()
    who = cell.paragraphs[0]
    pset(who, after=8)
    _tg_run(who, title, 9, True, TG_WHO)

    i = 0
    first_block = True
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip().startswith("…") or line.strip().startswith("..."):
            p = cell.add_paragraph()
            pset(p, after=4, before=6)
            _tg_run(p, line.strip(), 9, False, TG_MUTED)
            i += 1
            continue
        if _is_header(line):
            mark, rest = _split_emoji(line.strip())
            p = cell.add_paragraph()
            pset(p, after=4, before=0 if first_block else 10)
            first_block = False
            if mark:
                _tg_run(p, mark + " ", 12, True, TG_LINE, emoji=True)
            _tg_run(p, rest, 11, True, TG_LINE)
            i += 1
            continue
        if _is_status(line):
            mark, rest = _split_emoji(line.strip())
            p = cell.add_paragraph()
            pset(p, after=2, before=10)
            if mark:
                _tg_run(p, mark + " ", 11, True, TG_OK, emoji=True)
            _tg_run(p, rest, 11, True, TG_OK)
            i += 1
            continue
        batch = []
        while i < len(lines):
            cur = lines[i]
            if not cur.strip() or _is_header(cur) or _is_status(cur):
                break
            if cur.strip().startswith("…") or cur.strip().startswith("..."):
                break
            parsed = _split_metric(cur)
            if parsed:
                batch.append(parsed)
                i += 1
                continue
            break
        if batch:
            _add_metric_table(cell, batch)
            continue
        p = cell.add_paragraph()
        pset(p, after=3)
        text = line.strip()
        if text.startswith("✅"):
            _tg_run(p, text, 10.5, True, TG_OK)
        else:
            _tg_run(p, text, 10.5, False, TG_LINE)
        i += 1
    doc.add_paragraph()
